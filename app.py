import base64
import io
import os
import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ==============================================================================
# 1. إعدادات الصفحة الأولية (تستدعى مرة واحدة فقط في البداية)
# ==============================================================================
st.set_page_config(
    page_title="النظام الضريبي المتكامل - طنطا أول",
    page_icon="🏢",
    layout="wide",
)

# ==============================================================================
# 2. الثوابت ومسارات الملفات
# ==============================================================================
SIMILAR_PRICES_DB_PATH = "similar_prices_db.xlsx"
TAX_PAYERS_DB_PATH = "tax_payers_db.xlsx"
TANTA_1_DB_PATH = "قاعده بيانات طنطا اول.xlsx"
MEMO_FILE_PATH = "tax_memo_doc"

# ==============================================================================
# 3. إدارة جلسة العمل (Session State)
# ==============================================================================
if "users" not in st.session_state:
    st.session_state.users = {"ahmedh321": "Ahmed23"}
if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
if "username" not in st.session_state:
    st.session_state.username = ""
if "contracts_data" not in st.session_state:
    st.session_state.contracts_data = []
if "activities_data" not in st.session_state:
    st.session_state.activities_data = []

# حقول التعبئة التلقائية للأنشطة
if "auto_act_name" not in st.session_state:
    st.session_state.auto_act_name = ""
if "auto_act_reg" not in st.session_state:
    st.session_state.auto_act_reg = ""
if "auto_act_file" not in st.session_state:
    st.session_state.auto_act_file = ""
if "auto_act_addr" not in st.session_state:
    st.session_state.auto_act_addr = ""
if "auto_act_code" not in st.session_state:
    st.session_state.auto_act_code = ""

# حقول التعبئة التلقائية للعقود
if "c_seller_name" not in st.session_state:
    st.session_state.c_seller_name = ""
if "c_seller_addr" not in st.session_state:
    st.session_state.c_seller_addr = ""
if "c_seller_reg" not in st.session_state:
    st.session_state.c_seller_reg = ""
if "c_seller_file" not in st.session_state:
    st.session_state.c_seller_file = ""

# ==============================================================================
# 4. الدوال المساعدة والحسابية (Word & Helper Functions)
# ==============================================================================
def add_page_border(doc):
    """إضافة إطار كلاسيكي لصفحات Word."""
    sec_pr = doc.sections[0]._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for border_name in ["top", "left", "bottom", "right"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "24")
        border.set(qn("w:color"), "003366")
        pg_borders.append(border)
    sec_pr.append(pg_borders)

def calculate_tax_payer_tax(amount):
    """حساب الضريبة وفق الشرائح التصاعدية والتصرفات العقارية."""
    if amount < 1:
        return 0.0
    elif 1 <= amount <= 249999:
        return 1000.0
    elif 250000 <= amount <= 499999:
        return 2500.0
    elif 500000 <= amount <= 999999:
        return 5000.0
    elif 1000000 <= amount <= 1999999:
        return amount * 0.005
    elif 2000000 <= amount <= 2999999:
        return amount * 0.0075
    else:
        return amount * 0.01

def load_excel_db(path):
    """قراءة ملفات Excel بأمان."""
    if os.path.exists(path):
        try:
            df = pd.read_excel(path)
            df.columns = df.columns.str.strip()
            return df
        except Exception:
            return None
    return None

def extract_field_from_df(df_row, keywords, default_val=""):
    """البحث المرن عن قيمة العمود في السجل بالكلمات المفتاحية."""
    for col in df_row.index:
        for kw in keywords:
            if kw in str(col):
                val = df_row[col]
                return str(val) if pd.notna(val) else default_val
    return default_val

# --- إنشاء مذكرات وبرديات Word ---
def generate_activity_memo_docx(act_data):
    """إنشاء مذكرة فحص الأنشطة الرسمية (.docx)."""
    doc = Document()
    add_page_border(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_head.add_run("مصلحة الضرائب المصرية\nمأمورية طنطا أول\nمذكرة فحص نشاط\n")
    r.bold = True
    r.font.size = Pt(16)

    tbl = doc.add_table(rows=6, cols=2)
    tbl.style = "Table Grid"

    info_pairs = [
        ("اسم الممول", act_data.get("اسم الممول", "")),
        ("رقم التسجيل الضريبي", act_data.get("رقم التسجيل", "")),
        ("رقم الملف الضريبي", act_data.get("رقم الملف", "")),
        ("عنوان النشاط / الممول", act_data.get("العنوان", "")),
        ("نوع النشاط", act_data.get("نوع النشاط", "")),
        ("الكيان القانوني", act_data.get("الكيان القانوني", "")),
    ]

    for idx, (label, val) in enumerate(info_pairs):
        row_cells = tbl.rows[idx].cells
        row_cells[1].text = f" {label}"
        row_cells[1].paragraphs[0].runs[0].bold = True
        row_cells[0].text = f" {val}"

    doc.add_paragraph("\n")

    p1 = doc.add_paragraph()
    r1 = p1.add_run("1. المحاسبة السابقة:")
    r1.bold = True
    r1.font.size = Pt(13)
    doc.add_paragraph(act_data.get("المحاسبة السابقة", ""))

    p2 = doc.add_paragraph()
    r2 = p2.add_run("2. بيان إقرارات الممول المقدمة:")
    r2.bold = True
    r2.font.size = Pt(13)

    t_dec = doc.add_table(rows=2, cols=4)
    t_dec.style = "Table Grid"
    headers = ["السنة", "رقم الأعمال", "صافي الربح", "الضريبة"]
    for i, h in enumerate(headers):
        t_dec.rows[0].cells[i].text = h
        t_dec.rows[0].cells[i].paragraphs[0].runs[0].bold = True

    t_dec.rows[1].cells[0].text = str(act_data.get("سنة الإقرار", ""))
    t_dec.rows[1].cells[1].text = f"{float(act_data.get('رقم الأعمال', 0)):,.2f} ج.م"
    t_dec.rows[1].cells[2].text = f"{float(act_data.get('صافي الربح', 0)):,.2f} ج.م"
    t_dec.rows[1].cells[3].text = f"{float(act_data.get('ضريبة الإقرار', 0)):,.2f} ج.م"

    doc.add_paragraph("\n")

    p3 = doc.add_paragraph()
    r3 = p3.add_run("3. الإخطارات المقدمة من الممول:")
    r3.bold = True
    r3.font.size = Pt(13)
    doc.add_paragraph(act_data.get("تفاصيل الإخطارات", "لا يوجد إخطارات مقدمة."))

    p4 = doc.add_paragraph()
    r4 = p4.add_run("4. نتائج محاضر الاطلاع على الجهات:")
    r4.bold = True
    r4.font.size = Pt(13)

    t_auth = doc.add_table(rows=4, cols=2)
    t_auth.style = "Table Grid"
    auth_data = [
        ("الخصم والتحصيل", act_data.get("جهة الخصم والتحصيل", "لا يوجد")),
        ("القيمة المضافة", act_data.get("جهة القيمة المضافة", "لا يوجد")),
        ("الجمارك", act_data.get("جهة الجمارك", "لا يوجد")),
        ("الفاتورة الإلكترونية", act_data.get("جهة الفاتورة الإلكترونية", "لا يوجد")),
    ]
    for idx, (label, val) in enumerate(auth_data):
        row_cells = t_auth.rows[idx].cells
        row_cells[1].text = f" {label}"
        row_cells[1].paragraphs[0].runs[0].bold = True
        row_cells[0].text = f" {val}"

    doc.add_paragraph("\n")

    p5 = doc.add_paragraph()
    r5 = p5.add_run("5. محاضر الأعمال:")
    r5.bold = True
    r5.font.size = Pt(13)
    doc.add_paragraph(act_data.get("محاضر الأعمال", ""))

    p6 = doc.add_paragraph()
    r6 = p6.add_run("6. الفحص والمحاسبة:")
    r6.bold = True
    r6.font.size = Pt(13)
    doc.add_paragraph(act_data.get("تفاصيل الفحص والمحاسبة", ""))

    doc.add_paragraph("\n\n")
    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.autofit = False

    cell_auditor = t_sig.rows[0].cells[0]
    cell_reviewer = t_sig.rows[0].cells[1]

    cell_auditor.paragraphs[0].add_run("المأمور الفاحص:\n\n....................")
    cell_auditor.paragraphs[0].runs[0].bold = True
    cell_auditor.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    cell_reviewer.paragraphs[0].add_run("المراجع:\n\n....................")
    cell_reviewer.paragraphs[0].runs[0].bold = True
    cell_reviewer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def generate_tax_memo_docx(memo_data):
    """إنشاء مذكرة الفحص الرسمية للعقود (.docx)."""
    doc = Document()
    add_page_border(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_head.add_run("مصلحة الضرائب المصرية\nمذكرة الفحص - تصرفات عقارية\n")
    r.bold = True
    r.font.size = Pt(16)

    tbl = doc.add_table(rows=6, cols=2)
    tbl.style = "Table Grid"

    info_pairs = [
        ("اسم الممول", memo_data.get("اسم الممول", "")),
        ("رقم التسجيل الضريبي", memo_data.get("رقم التسجيل الضريبي", "")),
        ("رقم الملف الضريبي", memo_data.get("رقم الملف", "")),
        ("الكيان القانوني", memo_data.get("الكيان القانوني", "شخص طبيعي")),
        ("السنوات محل الفحص", memo_data.get("السنوات", "")),
        ("العنوان", memo_data.get("العنوان", "")),
    ]

    for idx, (label, val) in enumerate(info_pairs):
        row_cells = tbl.rows[idx].cells
        row_cells[1].text = f" {label}"
        row_cells[1].paragraphs[0].runs[0].bold = True
        row_cells[0].text = f" {val}"

    doc.add_paragraph("\n")

    sections = [
        ("مقدمة:", memo_data.get("المقدمة", "بناءً على التصرفات العقارية الواردة...")),
        ("بيان إقرارات الممول:", memo_data.get("إقرارات الممول", "لم يقدم الممول إقرارات عن العملية.")),
        ("محاضر الأعمال:", memo_data.get("محاضر الأعمال", "تم استيفاء البيانات والمتابعة.")),
    ]

    for sec_title, sec_text in sections:
        p_s = doc.add_paragraph()
        r_s = p_s.add_run(sec_title)
        r_s.bold = True
        r_s.font.size = Pt(13)
        doc.add_paragraph(sec_text)

    p_f = doc.add_paragraph()
    r_f = p_f.add_run("الفحص والمحاسبة:")
    r_f.bold = True
    r_f.font.size = Pt(13)

    legal_text = (
        "استناداً إلى معطيات الفحص وحالات المثل المتاحة والبيانات المسجلة، "
        f"ونظراً للتصرف العقاري المحدد، تم احتساب الضريبة المستحقة وفقاً للقوانين واللوائح التنفيذية."
    )
    doc.add_paragraph(legal_text)

    p_b = doc.add_paragraph()
    r_b = p_b.add_run("أسس المحاسبة والضريبة المستحقة:")
    r_b.bold = True

    t_basis = doc.add_table(rows=0, cols=2)
    t_basis.style = "Table Grid"

    basis_items = [
        ("نوع المحاسبة", memo_data.get("نوع المحاسبة", "تصرفات عقارية")),
        ("إجمالي رقم الأعمال / ثمن البيع", f"{float(memo_data.get('سعر البيع', 0)):,.2f} ج.م"),
        ("طريقة الحساب", memo_data.get("طريقة الحساب", "")),
        ("الضريبة المستحقة", f"{float(memo_data.get('الضريبة', 0)):,.2f} ج.م"),
    ]

    for l, v in basis_items:
        r_c = t_basis.add_row().cells
        r_c[1].text = f" {l}"
        r_c[1].paragraphs[0].runs[0].bold = True
        r_c[0].text = f" {v}"

    doc.add_paragraph("\n\n")
    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.autofit = False

    cell_auditor = t_sig.rows[0].cells[0]
    cell_reviewer = t_sig.rows[0].cells[1]

    cell_auditor.paragraphs[0].add_run("المأمور الفاحص:\n\n....................")
    cell_auditor.paragraphs[0].runs[0].bold = True
    cell_auditor.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    cell_reviewer.paragraphs[0].add_run("المراجع:\n\n....................")
    cell_reviewer.paragraphs[0].runs[0].bold = True
    cell_reviewer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def create_word_report(data):
    """إنشاء تقرير Word إجمالي لكافة العقود المفحوصة."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("تقرير فحص التصرفات العقارية والضريبة المحسوبة")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_paragraph("=" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER

    for item in data:
        p_header = doc.add_paragraph()
        r_h = p_header.add_run(f"📄 المستند: {item.get('اسم الملف', '')} ({item.get('نوع الحساب', '')})")
        r_h.bold = True
        r_h.font.size = Pt(14)

        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"

        for key, val in item.items():
            if key in ["اسم الملف"]:
                continue
            row_cells = table.add_row().cells
            row_cells[0].text = str(val)
            row_cells[1].text = str(key)
            row_cells[1].paragraphs[0].runs[0].bold = True

        doc.add_paragraph("\n" + "-" * 40 + "\n")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==============================================================================
# 5. شاشة تسجيل الدخول الموحدة
# ==============================================================================
if not st.session_state.logged_in:
    st.title("🔐 تسجيل الدخول للنظام الضريبي المتكامل")
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        with st.form("login_form"):
            user_input = st.text_input("اسم المستخدم:")
            pass_input = st.text_input("كلمة المرور:", type="password")
            submit_login = st.form_submit_button("تسجيل الدخول")

            if submit_login:
                if user_input in st.session_state.users and st.session_state.users[user_input] == pass_input:
                    st.session_state.logged_in = True
                    st.session_state.username = user_input
                    st.rerun()
                else:
                    st.error("اسم المستخدم أو كلمة المرور غير صحيحة!")
    st.stop()

# ==============================================================================
# 6. القائمة الجانبية (Sidebar)
# ==============================================================================
with st.sidebar:
    st.write(f"👤 **المستخدم الحالي:** `{st.session_state.username}`")
    if st.button("🚪 تسجيل الخروج"):
        st.session_state.logged_in = False
        st.session_state.username = ""
        st.rerun()

    st.markdown("---")

    selected_system = st.radio(
        "🔄 **اختر النظام العملياتي:**",
        ["🏭 نظام محاسبة الأنشطة", "📜 نظام مراجعة العقود والتصرفات"],
        index=0,
    )

    st.markdown("---")
    st.header("⚙️ إدارة قواعد البيانات")

    if selected_system == "🏭 نظام محاسبة الأنشطة":
        st.subheader("📊 قاعدة بيانات طنطا أول")
        uploaded_tanta = st.file_uploader(
            "رفع / تحديث قاعدة بيانات طنطا أول:",
            type=["xlsx", "xls"],
            key="upload_tanta",
        )
        if uploaded_tanta:
            with open(TANTA_1_DB_PATH, "wb") as f:
                f.write(uploaded_tanta.getbuffer())
            st.success("تم تحديث قاعدة بيانات طنطا أول بنجاح!")

        if os.path.exists(TANTA_1_DB_PATH):
            with st.expander("👁️ معاينة بيانات طنطا أول"):
                df_preview = load_excel_db(TANTA_1_DB_PATH)
                if df_preview is not None:
                    st.dataframe(df_preview.head(10), use_container_width=True)

    else:
        st.subheader("1️⃣ شيت حالات المثل")
        uploaded_similar = st.file_uploader(
            "رفع / تحديث شيت حالات المثل:", type=["xlsx", "xls"], key="upload_similar"
        )
        if uploaded_similar:
            with open(SIMILAR_PRICES_DB_PATH, "wb") as f:
                f.write(uploaded_similar.getbuffer())
            st.success("تم حفظ شيت حالات المثل بنجاح!")

        st.subheader("2️⃣ شيت الممولين المسجلين")
        uploaded_taxpayers = st.file_uploader(
            "رفع / تحديث بيانات الممولين المسجلين:",
            type=["xlsx", "xls"],
            key="upload_taxpayers",
        )
        if uploaded_taxpayers:
            with open(TAX_PAYERS_DB_PATH, "wb") as f:
                f.write(uploaded_taxpayers.getbuffer())
            st.success("تم حفظ شيت الممولين المسجلين بنجاح!")

    st.markdown("---")

    if st.session_state.username == "ahmedh321":
        st.header("👑 لوحة تحكم المسؤول (Admin)")

        with st.expander("📄 رفع/تحديث قالب مذكرة الفحص"):
            uploaded_memo = st.file_uploader(
                "اختر ملف المذكرة الأساسية (DOC, DOCX, PDF):",
                type=["pdf", "doc", "docx"],
                key="admin_memo_upload",
            )
            if uploaded_memo:
                ext = uploaded_memo.name.split(".")[-1]
                file_save_path = f"{MEMO_FILE_PATH}.{ext}"
                with open(file_save_path, "wb") as f:
                    f.write(uploaded_memo.getbuffer())
                st.success("تم حفظ المذكرة وتفعيل ربطها بنجاح!")

        with st.expander("➕ إضافة مستخدم جديد"):
            new_user = st.text_input("اسم المستخدم الجديد:")
            new_pass = st.text_input("كلمة المرور الجديدة:", type="password")
            if st.button("حفظ المستخدم"):
                if new_user and new_pass:
                    st.session_state.users[new_user] = new_pass
                    st.success(f"تم إضافة المستخدم '{new_user}' بنجاح!")

# ==============================================================================
# 7. النظام الأول: نظام محاسبة الأنشطة
# ==============================================================================
if selected_system == "🏭 نظام محاسبة الأنشطة":
    st.title("🏭 نظام محاسبة الأنشطة - مأمورية طنطا أول")

    col_left_act, col_right_act = st.columns([1, 1])

    with col_left_act:
        st.subheader("📄 رفع وتصفح مذكرات الفحص والمستندات")
        act_uploaded_file = st.file_uploader(
            "اختر مستند الفحص (PDF, صورة, DOCX):",
            type=["pdf", "png", "jpg", "jpeg", "docx"],
            key="act_file_upload",
        )

        if act_uploaded_file:
            if act_uploaded_file.type.startswith("image/"):
                st.image(act_uploaded_file, use_container_width=True)
            elif act_uploaded_file.type == "application/pdf":
                bytes_data = act_uploaded_file.getvalue()
                base64_pdf = base64.b64encode(bytes_data).decode("utf-8")
                pdf_display = f'<iframe src="data:application/pdf;base64,{base64_pdf}" width="100%" height="650" type="application/pdf"></iframe>'
                st.markdown(pdf_display, unsafe_allow_html=True)
            else:
                st.success(f"تم رفع المستند: {act_uploaded_file.name}")

    with col_right_act:
        st.subheader("📝 بيانات الفحص للأنشطة")

        def on_search_change():
            query = st.session_state.act_search_input.strip()
            tanta_df = load_excel_db(TANTA_1_DB_PATH)

            if tanta_df is not None and query:
                match = None
                for col in ["رقم التسجيل", "رقم الملف"]:
                    if col in tanta_df.columns:
                        m = tanta_df[tanta_df[col].astype(str).str.strip() == query]
                        if not m.empty:
                            match = m.iloc[0]
                            break

                if match is not None:
                    st.session_state.auto_act_name = str(match.get("اسم الممول", ""))
                    st.session_state.auto_act_reg = str(match.get("رقم التسجيل", ""))
                    st.session_state.auto_act_file = str(match.get("رقم الملف", ""))
                    st.session_state.auto_act_addr = str(match.get("العنوان", ""))
                    st.session_state.auto_act_code = str(match.get("كود النشاط", ""))
                    st.toast("✨ تم جلب واستيراد بيانات الممول من قاعدة البيانات بنجاح!")
                else:
                    st.toast("⚠️ لم يتم العثور على سجل مطابق برقم التسجيل/الملف.")

        st.text_input(
            "🔍 بحث برقم التسجيل أو رقم الملف (قاعدة بيانات طنطا أول):",
            key="act_search_input",
            on_change=on_search_change,
        )

        with st.form("activity_inspection_form"):
            st.markdown("##### 👤 البيانات الأساسية للممول والنشاط")
            act_taxpayer_name = st.text_input("اسم الممول:", value=st.session_state.auto_act_name)

            col_a1, col_a2 = st.columns(2)
            with col_a1:
                act_reg_num = st.text_input("رقم التسجيل الضريبي:", value=st.session_state.auto_act_reg)
                act_legal_entity = st.text_input("الكيان القانوني:", value="شخص طبيعي")
            with col_a2:
                act_file_num = st.text_input("رقم الملف الضريبي:", value=st.session_state.auto_act_file)
                act_type = st.text_input(
                    "نوع النشاط / كود النشاط:",
                    value=(f"نشاط تجاري (كود: {st.session_state.auto_act_code})" if st.session_state.auto_act_code else "نشاط تجاري"),
                )

            act_address = st.text_input("عنوان النشاط / الممول:", value=st.session_state.auto_act_addr)

            st.markdown("---")
            st.markdown("##### ⏳ 1. المحاسبة السابقة")
            prev_account_type = st.radio("وضع المحاسبة السابقة:", ["تمت المحاسبة حتى سنة", "الملف مستجد"], horizontal=True)

            if prev_account_type == "تمت المحاسبة حتى سنة":
                audit_until_year = st.selectbox("تمت المحاسبة حتى سنة:", options=[str(y) for y in range(2015, 2027)], index=7)
                prev_account_text = f"تمت محاسبة الممول حتى سنة {audit_until_year}."
            else:
                start_act_date = st.date_input("تاريخ بداية النشاط:")
                prev_account_text = f"الملف مستجد وتاريخ بداية نشاطه في: {start_act_date}."

            st.markdown("---")
            st.markdown("##### 📊 2. بيان الإقرارات المقدمة")
            col_dec1, col_dec2, col_dec3, col_dec4 = st.columns(4)
            with col_dec1:
                dec_year = st.selectbox("السنة:", [str(y) for y in range(2020, 2027)], index=5)
            with col_dec2:
                dec_turnover = st.number_input("رقم الأعمال (ج.م):", min_value=0.0, step=5000.0)
            with col_dec3:
                dec_net_profit = st.number_input("صافي الربح (ج.م):", min_value=0.0, step=1000.0)
            with col_dec4:
                dec_tax = st.number_input("الضريبة (ج.م):", min_value=0.0, step=500.0)

            st.markdown("---")
            st.markdown("##### 📩 3. الإخطارات المقدمة من الممول")
            has_notifications = st.radio("هل يوجد إخطارات مقدمة من الممول؟", ["لا", "نعم"], horizontal=True)
            if has_notifications == "نعم":
                notification_summary = st.text_area("ملخص الإخطار المقدم من الممول:", value="قدم الممول إخطاراً بموجب النموذج الضريبي يتضمن...")
            else:
                notification_summary = "لم يتقدم الممول بأي إخطارات توقف أو تغيير."

            st.markdown("---")
            st.markdown("##### 🔍 4. محاضر الاطلاع على الجهات")
            col_g1, col_g2 = st.columns(2)
            with col_g1:
                res_discount = st.text_input("الخصم والتحصيل (النتيجة):", value="لا يوجد تعاملات")
                res_vat = st.text_input("القيمة المضافة (النتيجة):", value="مسجل بالقيمة المضافة")
            with col_g2:
                res_customs = st.text_input("الجمارك (النتيجة):", value="لا يوجد رسوم جمركية")
                res_einvoice = st.text_input("الفاتورة الإلكترونية (النتيجة):", value="ملتزم منظومة الفواتير")

            st.markdown("---")
            st.markdown("##### 📝 5. محاضر الأعمال")
            work_minutes_text = st.text_area(
                "تفاصيل محاضر الأعمال:",
                value="تم إخطار الممول بموجب النماذج الضريبية وتحديد المواعيد للمناقشة والفحص.",
                height=80,
            )

            st.markdown("---")
            st.markdown("##### ⚖️ 6. الفحص والمحاسبة")
            audit_assessment_text = st.text_area(
                "تفاصيل الفحص والمحاسبة بالأسس والتقديرات:",
                value="استناداً إلى معطيات الفحص وحالات المثل المتاحة والبيانات المسجلة، قدر رقم الأعمال التقديري...",
                height=100,
            )

            submit_act_btn = st.form_submit_button("💾 حفظ بيانات الفحص")

            if submit_act_btn:
                act_record = {
                    "اسم الممول": act_taxpayer_name,
                    "رقم التسجيل": act_reg_num,
                    "رقم الملف": act_file_num,
                    "العنوان": act_address,
                    "الكيان القانوني": act_legal_entity,
                    "نوع النشاط": act_type,
                    "المحاسبة السابقة": prev_account_text,
                    "سنة الإقرار": dec_year,
                    "رقم الأعمال": dec_turnover,
                    "صافي الربح": dec_net_profit,
                    "ضريبة الإقرار": dec_tax,
                    "هل يوجد إخطارات": has_notifications,
                    "تفاصيل الإخطارات": notification_summary,
                    "جهة الخصم والتحصيل": res_discount,
                    "جهة القيمة المضافة": res_vat,
                    "جهة الجمارك": res_customs,
                    "جهة الفاتورة الإلكترونية": res_einvoice,
                    "محاضر الأعمال": work_minutes_text,
                    "تفاصيل الفحص والمحاسبة": audit_assessment_text,
                    "الفاحص": st.session_state.username,
                }
                st.session_state.activities_data.append(act_record)
                st.success("تم حفظ كافة بيانات النشاط بنجاح!")

        if st.session_state.activities_data:
            st.markdown("---")
            latest_act = st.session_state.activities_data[-1]
            act_doc_bytes = generate_activity_memo_docx(latest_act)

            st.download_button(
                label="🖨️ طباعة وتنزيل مذكرة فحص النشاط (.docx)",
                data=act_doc_bytes,
                file_name=f"مذكرة_فحص_نشاط_{latest_act.get('اسم الممول', 'الممول')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

# ==============================================================================
# 8. النظام الثاني: نظام مراجعة العقود والتصرفات العقارية
# ==============================================================================
elif selected_system == "📜 نظام مراجعة العقود والتصرفات":
    st.title("🏢 نظام مراجعة العقود وحساب الضريبة")

    with st.sidebar:
        uploaded_files = st.file_uploader(
            "📁 رفع عقود البيع (PDF, JPG, PNG):",
            type=["jpg", "jpeg", "png", "pdf"],
            accept_multiple_files=True,
        )

    if uploaded_files:
        current_file = uploaded_files[0]
        st.info(f"📄 معالجة العقد: **{current_file.name}**")

        col_right, col_left = st.columns([1, 1])

        with col_right:
            st.subheader("👁️ معاينة المستند الأصلي")
            if current_file.type.startswith("image/"):
                st.image(current_file, use_container_width=True)
            elif current_file.type == "application/pdf":
                bytes_data = current_file.getvalue()
                base64_pdf = base64.b64encode(bytes_data).decode("utf-8")
                pdf_display = f'<iframe src="data:application/pdf;base64,{base64_pdf}" width="100%" height="700" type="application/pdf"></iframe>'
                st.markdown(pdf_display, unsafe_allow_html=True)

        with col_left:
            st.subheader("📝 بيانات العقد وحساب الضريبة")

            def on_contract_taxpayer_search():
                q = st.session_state.contract_tax_search.strip()
                taxpayers_df = load_excel_db(TAX_PAYERS_DB_PATH)
                if taxpayers_df is not None and q:
                    for col in ["رقم التسجيل", "تسجيل", "رقم الملف"]:
                        matches = [c for c in taxpayers_df.columns if col in str(c)]
                        if matches:
                            m = taxpayers_df[taxpayers_df[matches[0]].astype(str).str.strip() == q]
                            if not m.empty:
                                r = m.iloc[0]
                                st.session_state.c_seller_name = extract_field_from_df(r, ["اسم", "ممول"], "")
                                st.session_state.c_seller_addr = extract_field_from_df(r, ["عنوان", "موقع"], "")
                                st.session_state.c_seller_reg = extract_field_from_df(r, ["تسجيل", "رقم"], q)
                                st.session_state.c_seller_file = extract_field_from_df(r, ["ملف"], "")
                                st.toast("✨ تم جلب بيانات الممول المسجل بالعقود بنجاح!")
                                break

            st.text_input(
                "🔍 بحث برقم التسجيل للممولين المسجلين:",
                key="contract_tax_search",
                on_change=on_contract_taxpayer_search,
            )

            calc_method = st.radio("طريقة حساب الضريبة:", ["نسبة ثابته (2.5%)", "وفق الشرائح التصاعدية"], horizontal=True)

            with st.form("contract_entry_form"):
                seller_name = st.text_input("اسم البائع / الممول:", value=st.session_state.c_seller_name)
                seller_id = st.text_input("الرقم القومي للبائع (14 رقم):")
                property_address = st.text_area("العنوان التفصيلي للعقار:", value=st.session_state.c_seller_addr)
                sale_price = st.number_input("ثمن البيع بالكتّاب (ج.م):", min_value=0.0, step=1000.0)

                # حساب الضريبة وفق الاختيار
                if calc_method == "نسبة ثابته (2.5%)":
                    tax_amount = sale_price * 0.025
                else:
                    tax_amount = calculate_tax_payer_tax(sale_price)

                st.write(f"💰 **الضريبة المستحقة:** `{tax_amount:,.2f} ج.م`")
                save_btn = st.form_submit_button("💾 حفظ بيانات العقد")

                if save_btn:
                    contract_data_dict = {
                        "اسم الملف": current_file.name,
                        "اسم البائع/الممول": seller_name,
                        "رقم التسجيل الضريبي": st.session_state.c_seller_reg,
                        "رقم الملف": st.session_state.c_seller_file,
                        "الرقم القومي": seller_id,
                        "العنوان": property_address,
                        "سعر البيع": sale_price,
                        "طريقة الحساب": calc_method,
                        "الضريبة": tax_amount,
                    }
                    st.session_state.contracts_data.append(contract_data_dict)
                    st.success("تم حفظ بيانات العقد بنجاح!")

            if st.session_state.contracts_data:
                st.markdown("---")
                latest_contract = st.session_state.contracts_data[-1]
                contract_memo_bytes = generate_tax_memo_docx({
                    "اسم الممول": latest_contract.get("اسم البائع/الممول"),
                    "رقم التسجيل الضريبي": latest_contract.get("رقم التسجيل الضريبي"),
                    "رقم الملف": latest_contract.get("رقم الملف"),
                    "الكيان القانوني": "شخص طبيعي",
                    "السنوات": "تصرف عقاري",
                    "العنوان": latest_contract.get("العنوان"),
                    "المقدمة": "فحص تصرف عقاري بموجب عقد بيع مثبت.",
                    "إقرارات الممول": "عدم تقديم إقرار عن التصرف.",
                    "محاضر الأعمال": "تم مراجعة العقد وحساب الضريبة.",
                    "نوع المحاسبة": "تصرفات عقارية",
                    "سعر البيع": latest_contract.get("سعر البيع"),
                    "طريقة الحساب": latest_contract.get("طريقة الحساب"),
                    "الضريبة": latest_contract.get("الضريبة"),
                })

                st.download_button(
                    label="🖨️ تنزيل مذكرة فحص التصرف العقاري (.docx)",
                    data=contract_memo_bytes,
                    file_name=f"مذكرة_تصرف_عقاري_{latest_contract.get('اسم البائع/الممول', 'ممول')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
    else:
        st.info("👈 يرجى رفع عقود البيع من القائمة الجانبية لبدء العمل.")

# ==============================================================================
# 9. عرض السجلات والتصدير الإجمالي
# ==============================================================================
if st.session_state.activities_data or st.session_state.contracts_data:
    st.markdown("---")
    st.subheader("📊 إجمالي السجلات المفحوصة")

    if selected_system == "🏭 نظام محاسبة الأنشطة" and st.session_state.activities_data:
        st.dataframe(pd.DataFrame(st.session_state.activities_data), use_container_width=True)

    elif selected_system == "📜 نظام مراجعة العقود والتصرفات" and st.session_state.contracts_data:
        st.dataframe(pd.DataFrame(st.session_state.contracts_data), use_container_width=True)

        report_word_bytes = create_word_report(st.session_state.contracts_data)
        st.download_button(
            label="📄 تنزيل تقرير Word الشامل لجميع العقود",
            data=report_word_bytes,
            file_name="تقرير_فحص_التصرفات_العقارية_الإجمالي.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )